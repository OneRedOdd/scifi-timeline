# Build paper.docx from paper.md (pandoc) and style it; then convert to PDF with LibreOffice.
import pypandoc, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
pypandoc.convert_file('paper.md','docx',outputfile='paper_raw.docx',extra_args=['--resource-path=.:..'])
d=Document('paper_raw.docx')
for s in d.sections:
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Cm(2.2)
st=d.styles; names=[x.name for x in st]
for name in ['Normal','Body Text','First Paragraph','Compact']:
    if name in names:
        f=st[name].font; f.name='Calibri'; f.size=Pt(11)
for name,size in [('Title',22),('Heading 1',15),('Heading 2',12.5)]:
    s=st[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0x1f,0x3a,0x5f)
def shade(c,fill):
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
for tbl in d.tables:
    tblPr=tbl._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'8B97B0'); b.append(el)
    tblPr.append(b)
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    w=OxmlElement('w:tblW'); w.set(qn('w:w'),'5000'); w.set(qn('w:type'),'pct'); tblPr.append(w)
    for i,row in enumerate(tbl.rows):
        if i==0:
            trPr=row._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                for r in p.runs:
                    r.font.size=Pt(9)
                    if i==0: r.font.bold=True; r.font.color.rgb=RGBColor(0xff,0xff,0xff)
            if i==0: shade(c,'1F3A5F')
            elif i%2==0: shade(c,'EEF2F8')
for p in d.paragraphs:
    if p.style.name in ('Image Caption','Table Caption'):
        for r in p.runs: r.font.size=Pt(9)
d.save('paper.docx')
import os; os.remove('paper_raw.docx')
subprocess.run(['soffice','--headless','--convert-to','pdf','paper.docx'],check=True,capture_output=True)
